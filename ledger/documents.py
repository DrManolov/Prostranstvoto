"""Render by editing copies of the July OOXML package, retaining its images and styles."""
from copy import deepcopy
from datetime import date,datetime,timezone as dt_timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import uuid
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db.models import Q
from .models import GeneratedDocument, Official, SourceDocument
from .services import audit, period_totals, serial, writer

SECTIONS={'plan':(0,22),'leader_report':(22,58),'dean_report':(58,90),'external_order':(91,120),'staff_order':(121,152)}
NAMES={'plan':'План-сметка','leader_report':'Доклад до декана','dean_report':'Доклад до ректора',
       'external_order':'Заповед за нещатни','staff_order':'Заповед за щатни','invoice_order':'Заповед за фактури и други разходи','combined':'Общ документ'}

def txt(node): return ''.join(t.text or '' for t in node.iter(qn('w:t')))

def set_text(node,value):
    if txt(node)==str(value): return
    old=txt(node)
    if node.tag==qn('w:tc') and str(value).endswith(' евро'):
        import re
        match=re.match(r'^\s*([\d\s\u00a0]+(?:[,.]\d+)?)\s*(евро|лв\.?)',old)
        if match:
            replace_text(node,match.group(1).strip(),str(value)[:-5].strip())
            if match.group(2)!='евро': replace_text(node,match.group(2),'евро')
            return
    nodes=list(node.iter(qn('w:t')))
    if nodes:
        nodes[0].text=str(value); nodes[0].set(qn('xml:space'),'preserve')
        for n in nodes[1:]: n.text=''
    else:
        p=node if node.tag==qn('w:p') else node.find(qn('w:p'))
        if p is None: p=OxmlElement('w:p'); node.append(p)
        r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=str(value); r.append(t); p.append(r)

def replace_text(node,old,new):
    # Replace across split Word runs while retaining formatting of untouched runs.
    parts=list(node.iter(qn('w:t'))); text=''.join(x.text or '' for x in parts)
    if old not in text or old==new: return
    import re
    positions=[m.start() for m in re.finditer(re.escape(old),text)]
    for start in reversed(positions):
        end=start+len(old); offset=0; first=True
        for part in parts:
            s=part.text or ''; a,b=offset,offset+len(s); offset=b
            if b<=start or a>=end: continue
            left=s[:max(0,start-a)]; right=s[max(0,end-a):] if end<b else ''
            part.text=left+(str(new) if first else '')+right; first=False

def euro(value): return f'{Decimal(value):,.2f}'.replace(',',' ').replace('.',',')+' евро'

def rows(table,data):
    originals=list(table.findall(qn('w:tr')))
    prototype=deepcopy(originals[-1])
    for r in originals: table.remove(r)
    for index,values in enumerate(data):
        style=originals[index] if len(data)==len(originals) else next((r for r in originals if txt(r.findall(qn('w:tc'))[0])==values[0]),originals[0])
        row=deepcopy(style)
        cells=row.findall(qn('w:tc'))
        for c,v in zip(cells,values): set_text(c,v)
        table.append(row)

def people_table(table,people,category):
    original=list(table.findall(qn('w:tr')))
    header=deepcopy(original[0]); template=deepcopy(original[-1])
    for row in original: table.remove(row)
    if category=='invoice':
        set_text(header.findall(qn('w:tc'))[1],'Плащания по фактури и други допустими разходи')
    table.append(header)
    for p in people:
        row=deepcopy(template)
        cells=row.findall(qn('w:tc'))
        for cell,value in zip(cells,['',p['name'],euro(p['amount'])]): set_text(cell,value)
        table.append(row)

def officials_for(period,day):
    if period.historical:
        text=period.source.text
        import re
        plan=text.split('ДО ДЕКАНА')[0]
        names=re.search(r'\((проф[^)]+)\)\s+\((проф[^)]+)\)',plan,re.I)
        leader=names.group(1) if names else 'проф. д.пс.н. Манол Манолов'
        dean=names.group(2) if names else 'проф. д-р Владимира Ангелова'
        return {'rector':'проф. д-р Румен Младенов','dean':dean,
                'leader':leader,'accountant':'Виеслав Пиларски'}
    result={}
    for role in ['rector','dean','leader','accountant']:
        names=list(Official.objects.filter(role=role,confirmed=True,valid_from__lte=day).filter(Q(valid_until__isnull=True)|Q(valid_until__gte=day)))
        if len(names)!=1: raise ValidationError(f'Потвърдете точно едно длъжностно лице за {role} към датата на документа в административните настройки.')
        result[role]=names[0].name
    return result

def render_section(template,key,period,data,officials,day):
    document=Document(template)
    document.core_properties.title=NAMES[key]+' — '+period.title
    document.core_properties.subject='Проект за административно утвърждаване'
    document.core_properties.created=document.core_properties.modified=datetime.now(dt_timezone.utc)
    nodes=list(document.element.body)
    start,end=SECTIONS['external_order' if key=='invoice_order' else key]
    selected=[deepcopy(n) for n in nodes[start:end]]
    for node in list(document.element.body):
        if node.tag!=qn('w:sectPr'): document.element.body.remove(node)
    body=document.element.body
    staff=[p for p in data.get('expense_people',[]) if p['category']=='staff']
    external=[p for p in data.get('expense_people',[]) if p['category'] in ('external','pensioner')]
    invoices=[p for p in data.get('expense_people',[]) if p['category'] in ('invoice','other','insurance')]
    tables=[n for n in selected if n.tag==qn('w:tbl')]
    role_correction=(not period.historical or period.discrepancies.filter(code='dean_signature_role',state='accepted').exists())
    if key=='plan':
        if period.historical:
            plan=[(v.label,euro(v.effective_eur)) for v in period.values.all()]
        else:
            plan=[(r['label'],euro(r['amount'])) for r in data['revenues']]
            rates=data.get('rule_rates',[])
            if len(rates)>1:
                for rule in rates:
                    plan.append((f"Финансов режим «{rule['name']}»: университет {rule['university']}%, факултет {rule['faculty']}%, ПФ+ {rule['center']}% — постъпления",euro(rule['gross'])))
            labels=[('total','Общо постъпления от провеждането на курсове'),('university','Отчисления за университета'),
                    ('faculty','Отчисления за Педагогически факултет'),('center','Отчисления за Пространството ПФ+'),
                    ('remaining','Остатък след отчисления'),('opening','Остатък от предходно плащане'),
                    ('available','Обща сума за разпределение'),('invoice','Материални разходи – плащания по фактура'),
                    ('staff','Трудови възнаграждения'),('insurance_staff','Социални и здравни осигуровки — щатни'),
                    ('external','Трудови възнаграждения на външни преподаватели'),('insurance_external','Социални и здравни осигуровки — външни'),
                    ('pensioner','Хонорари — пенсионери'),('insurance_pensioner','Социални и здравни осигуровки — пенсионери'),
                    ('other','Други разходи'),('insurance','Отделно начислени осигуровки'),('costs','Обща сума разходи'),('closing','Остатък за следващо плащане')]
            for k,label in labels:
                if not Decimal(data.get(k,'0')) and k not in ('total','university','faculty','center','remaining','opening','available','costs','closing'): continue
                if len(rates)==1 and k in ('university','faculty','center','remaining'):
                    rate=Decimal(rates[0][k]) if k!='remaining' else 100-sum(Decimal(rates[0][x]) for x in ('university','faculty','center'))
                    label+=' ('+format(rate.normalize(),'f').replace('.',',')+'%)'
                if k.startswith('insurance_'):
                    percents={r['percent'] for r in data.get('compensation_rates',[]) if r['category']==k[10:]}
                    if len(percents)==1: label+=' ('+next(iter(percents)).replace('.',',')+'%)'
                plan.append((label,euro(data.get(k,'0'))))
        rows(tables[0],plan)
    elif key in ('leader_report','dean_report'):
        people_table(tables[0],staff,'staff'); people_table(tables[1],external,'external')
        if invoices:
            clone=deepcopy(tables[1]); people_table(clone,invoices,'invoice')
            pos=selected.index(tables[1])+2; selected.insert(pos,clone)
            summary=deepcopy(next(n for n in selected if txt(n).startswith('СУМА ЗА ХОНОРАРИ:')))
            set_text(summary,'СУМА ЗА ПЛАЩАНИЯ ПО ФАКТУРИ И ДРУГИ РАЗХОДИ: '+euro(sum(Decimal(p['amount']) for p in invoices)))
            selected.insert(pos+1,summary)
    else:
        group=staff if key=='staff_order' else invoices if key=='invoice_order' else external
        people_table(tables[0],group,'invoice' if key=='invoice_order' else key)
    narrative=', '.join(dict.fromkeys(r['label'].split(',')[0] for r in data.get('revenues',[]))) or 'дейността на Центъра'
    for n in list(selected):
        if n.tag!=qn('w:p'): continue
        text=txt(n)
        if 'Въз основа на' in text:
            # Retain administrative framing and its exact paragraph/run properties.
            a=text.find('по курсове '); b=text.find('и такси от други курсове')
            if a>=0 and b>a: replace_text(n,text[a+len('по курсове '):b],narrative+' ')
        if 'КЪМ ПЕДАГОГИЧЕСКИЯ ФАКУЛТЕТ, МЕСЕЦ' in text:
            set_text(n,'КЪМ ПЕДАГОГИЧЕСКИЯ ФАКУЛТЕТ, '+period.title.upper())
        if 'Дата: 14.07.2026' in text: replace_text(n,'14.07.2026',day.strftime('%d.%m.%Y'))
        if text.startswith('ОСТАТЪК ЗА СЛЕДВАЩО ПЛАЩАНЕ:'): set_text(n,'ОСТАТЪК ЗА СЛЕДВАЩО ПЛАЩАНЕ: '+euro(data.get('closing','0')))
        if text.startswith('СУМА ЗА ВЪЗНАГРАЖДЕНИЯ:'):
            if not staff: selected.remove(n); continue
            set_text(n,'СУМА ЗА ВЪЗНАГРАЖДЕНИЯ: '+euro(sum(Decimal(p['amount']) for p in staff)))
        if text.startswith('СУМА ЗА ХОНОРАРИ:'):
            group=invoices if key=='invoice_order' else external
            if not group: selected.remove(n); continue
            set_text(n,('СУМА ЗА ПЛАЩАНИЯ: ' if key=='invoice_order' else 'СУМА ЗА ХОНОРАРИ: ')+euro(sum(Decimal(p['amount']) for p in group)))
        if text.strip()=='ЩАТНИ ПРЕПОДАВАТЕЛИ' and not staff: selected.remove(n); continue
        if role_correction and key=='dean_report' and text.strip()=='РЪКОВОДИТЕЛ НА КУРСОВЕ КЪМ ПЕДАГОГИЧЕСКИЯ ФАКУЛТЕТ':
            set_text(n,'ДЕКАН НА ПЕДАГОГИЧЕСКИЯ ФАКУЛТЕТ')
        if text.strip().startswith('НЕЩАТНИ ПРЕПОДАВАТЕЛИ'):
            if key=='invoice_order': set_text(n,'ПЛАЩАНИЯ ПО ФАКТУРИ И ДРУГИ РАЗХОДИ')
            elif not external: selected.remove(n); continue
        if key=='invoice_order' and text.strip()=='Да се изплатят следните възнаграждения:':
            set_text(n,'Да се извършат следните плащания:')
        replacements={'проф. д-р Румен Младенов':officials['rector'], 'ПРОФ. Д-Р РУМЕН МЛАДЕНОВ':officials['rector'].upper(),
           'проф. д-р Владимира Ангелова':officials['dean'],'ПРОФ. Д-Р ВЛАДИМИРА АНГЕЛОВА':officials['dean'].upper(),
           'проф. д.пс.н. Манол Манолов':officials['leader'],'ПРОФ. Д.ПС.Н. МАНОЛ НИКОЛОВ МАНОЛОВ':officials['leader'].upper(),
           'ПРОФ. Д.ПС.Н МАНОЛ МАНОЛОВ':officials['leader'].upper(),'Виеслав Пиларски':officials['accountant']}
        for old,new in replacements.items():
            if old!=new: replace_text(n,old,new)
    if key in ('leader_report','dean_report'):
        if not staff: selected.remove(tables[0])
        if not external: selected.remove(tables[1])
    for n in selected: body.insert(len(body)-1,n)
    # A draft label in the footer keeps the original body, header graphics, and margins.
    for section in document.sections:
        from docx.shared import Mm
        section.page_width=Mm(210); section.page_height=Mm(297)
        p=section.footer.paragraphs[0]
        p.text='ПРОЕКТ • '+period.title+' • '+day.strftime('%d.%m.%Y')
        for run in p.runs:
            from docx.shared import Pt
            run.font.size=Pt(8)
    return document

def to_pdf(docx_path):
    with tempfile.TemporaryDirectory(prefix='pfplus-lo-') as temp:
        command=[settings.SOFFICE,'-env:UserInstallation=file://'+temp+'/profile','--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(docx_path.parent),str(docx_path)]
        try: result=subprocess.run(command,capture_output=True,text=True,timeout=90)
        except (FileNotFoundError,subprocess.TimeoutExpired) as e:
            raise ValidationError('PDF конверторът LibreOffice е недостъпен. Използвайте Docker конфигурацията или инсталирайте LibreOffice Writer.') from e
        pdf=docx_path.with_suffix('.pdf')
        if result.returncode or not pdf.exists() or pdf.stat().st_size<500: raise ValidationError('Неуспешно PDF преобразуване. DOCX е запазен; проверете LibreOffice и шрифтовете.')
        return pdf

def generate(period,user,day=None):
    from django.utils import timezone
    writer(user); day=day or timezone.localdate()
    source=SourceDocument.objects.filter(path__contains='/01_').first()
    if source is None:
        raise ValidationError('Липсва оригиналният юлски DOCX шаблон. Добавете частните източници според data/README.md и изпълнете import_history.')
    template=settings.BASE_DIR/source.path
    if not template.is_file():
        raise ValidationError('Оригиналният DOCX шаблон липсва от частните файлове. Възстановете го от защитеното копие.')
    data=period.snapshot if period.state!='draft' else period_totals(period)
    officials=officials_for(period,day)
    if hashlib.sha256(template.read_bytes()).hexdigest()!=source.sha256: raise ValidationError('Контролната сума на шаблона не съвпада.')
    data=dict(data,officials=officials,document_date=str(day),template_sha256=source.sha256,period_revision=period.revision)
    data['discrepancy_decisions']=list(period.discrepancies.values('id','code','state','reason'))
    folder=settings.MEDIA_ROOT/'generated'/uuid.uuid4().hex
    folder.mkdir(parents=True,mode=0o700)
    keys=['plan','leader_report','dean_report']
    cats={p['category'] for p in data.get('expense_people',[]) if Decimal(p['amount'])}
    if cats & {'external','pensioner'}: keys.append('external_order')
    if 'staff' in cats: keys.append('staff_order')
    if cats & {'invoice','other','insurance'}: keys.append('invoice_order')
    combined=Document(template)
    combined.core_properties.title='Общ документ — '+period.title
    combined.core_properties.subject='Проект за административно утвърждаване'
    combined.core_properties.created=combined.core_properties.modified=datetime.now(dt_timezone.utc)
    for node in list(combined.element.body):
        if node.tag!=qn('w:sectPr'): combined.element.body.remove(node)
    files={}; pdf_errors=[]; drawing_id=1
    for i,key in enumerate(keys):
        doc=render_section(template,key,period,data,officials,day)
        target=folder/(key+'.docx'); doc.save(target)
        files[key+'.docx']=str(target.relative_to(settings.MEDIA_ROOT))
        section_nodes=[deepcopy(node) for node in doc.element.body if node.tag!=qn('w:sectPr')]
        if i and section_nodes:
            # Put the break on the first existing paragraph. A standalone break can
            # spill to an otherwise blank page when the preceding page is full.
            first=section_nodes[0]
            prop=first.find(qn('w:pPr'))
            if prop is None: prop=OxmlElement('w:pPr'); first.insert(0,prop)
            pb=OxmlElement('w:pageBreakBefore'); prop.append(pb)
        for node in section_nodes:
            for prop in node.iter(qn('wp:docPr')):
                prop.set('id',str(drawing_id)); drawing_id+=1
            combined.element.body.insert(len(combined.element.body)-1,node)
        try: files[key+'.pdf']=str(to_pdf(target).relative_to(settings.MEDIA_ROOT))
        except ValidationError as e: pdf_errors+=e.messages
    for section in combined.sections:
        from docx.shared import Mm
        section.page_width=Mm(210); section.page_height=Mm(297)
        section.footer.paragraphs[0].text='ПРОЕКТ • '+period.title+' • '+day.strftime('%d.%m.%Y')
    target=folder/'combined.docx'; combined.save(target)
    files['combined.docx']=str(target.relative_to(settings.MEDIA_ROOT))
    try: files['combined.pdf']=str(to_pdf(target).relative_to(settings.MEDIA_ROOT))
    except ValidationError as e: pdf_errors+=e.messages
    data['pdf_errors']=list(dict.fromkeys(pdf_errors))
    data['file_hashes']={name:hashlib.sha256((settings.MEDIA_ROOT/path).read_bytes()).hexdigest() for name,path in files.items()}
    (folder/'provenance.json').write_text(json.dumps(serial(data),ensure_ascii=False,indent=2))
    files['provenance.json']=str((folder/'provenance.json').relative_to(settings.MEDIA_ROOT))
    record=GeneratedDocument.objects.create(period=period,created_by=user,data=serial(data),files=files)
    audit(user,'Генериран проект на документи',record,after={'files':list(files),'period_revision':period.revision,'template_sha256':source.sha256})
    return record
