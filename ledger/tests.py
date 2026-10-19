from datetime import date,datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError,PermissionDenied
from django.core.management import call_command
from django.db import IntegrityError,transaction
from django.test import TestCase,override_settings
from django.utils import timezone
from .models import (Audit,CompensationRate,Course,Disbursement,Discrepancy,Expense,FinancialRule,
 HistoricalAdjustment,HistoricalValue,Participant,Payment,Period,SourceDocument)
from .money import allocate,eur,money
from .services import (budget,close_period,confirm_payment,dashboard_data,pay_expense,period_totals)

class FinanceTests(TestCase):
    def setUp(self):
        self.user=get_user_model().objects.create_superuser('admin','',None)
        self.viewer=get_user_model().objects.create_user('viewer')
        self.course=Course.objects.create(title='Тестов курс',start=timezone.make_aware(datetime(2027,1,15,10)),expected_count=10,expected_price=Decimal('90'),expected_cost=Decimal('100'))
        self.rule=FinancialRule.objects.create(name='Потвърдено правило',valid_from=date(2027,1,1),university=30,faculty=10,center=10,confirmed=True,accounting_reference='Счетоводство / тест')
        self.person=Participant.objects.create(course=self.course,name='Тестово лице',email='test@example.com',due=Decimal('90'))
    def payment(self,amount='45',**kwargs):
        return Payment.objects.create(date=date(2027,1,15),course=self.course,participant=self.person,payer='Тестово лице',
            amount=Decimal(amount),currency='EUR',reference=kwargs.pop('reference',f'ref-{Payment.objects.count()}'),created_by=self.user,**kwargs)
    def test_fixed_conversion_and_no_float(self):
        self.assertEqual(eur(Decimal('1955.83'),'BGN'),Decimal('1000.00'))
        self.assertEqual(eur(Decimal('2040'),'BGN'),Decimal('1043.04'))
        self.assertEqual(money('1.005'),Decimal('1.01'))
        with self.assertRaises(TypeError): money(1.005)
    def test_exact_allocation_rounding(self):
        shares,rest,delta=allocate('0.05',['30','10','10'])
        self.assertEqual(sum(shares)+rest,Decimal('0.05'))
        self.assertEqual(delta,Decimal('-0.02'))
    def test_three_acceptance_examples(self):
        examples=[('7146','505.43','3534.78','543.65'),('2412','543.65','1715.07','34.58'),('2880','34.58','719.26','755.32')]
        for total,opening,costs,closing in examples:
            shares,rest,_=allocate(total,['30','10','10'])
            self.assertEqual(rest+Decimal(opening)-Decimal(costs),Decimal(closing))
    def test_missing_private_template_has_actionable_error(self):
        from .documents import generate
        from .models import GeneratedDocument
        period=Period.objects.create(title='Януари 2027',start=date(2027,1,1),end=date(2027,1,31))
        with self.assertRaisesMessage(ValidationError,'Липсва оригиналният юлски DOCX шаблон'):
            generate(period,self.user)
        self.assertEqual(GeneratedDocument.objects.count(),0)
    def test_partial_payment_and_refund(self):
        p=self.payment(); self.assertEqual(self.person.paid,0)
        confirm_payment(p.pk,self.user); self.assertEqual(self.person.payment_status,'Частично платил')
        p2=self.payment(); confirm_payment(p2.pk,self.user)
        self.assertEqual(self.person.payment_status,'Платил')
        refund=self.payment('45',kind='refund',refund_of=p)
        confirm_payment(refund.pk,self.user)
        self.assertEqual(self.person.paid,Decimal('45'))
        refund2=self.payment('1',kind='refund',refund_of=p)
        with self.assertRaises(ValidationError): confirm_payment(refund2.pk,self.user)
    def test_confirmation_and_reference_deduplicate(self):
        p=self.payment(); confirm_payment(p.pk,self.user)
        with self.assertRaises(ValidationError): confirm_payment(p.pk,self.user)
        with self.assertRaises(ValidationError): self.payment(reference=p.reference)
        self.assertEqual(dashboard_data({})['total'],Decimal('45'))
        p.refresh_from_db(); p.amount=Decimal('100')
        with self.assertRaises(ValidationError): p.save()
    def test_refund_restores_all_allocated_cents(self):
        p=self.payment('0.05'); confirm_payment(p.pk,self.user); p.refresh_from_db()
        for _ in range(5):
            refund=self.payment('0.01',kind='refund',refund_of=p)
            confirm_payment(refund.pk,self.user)
        data=dashboard_data({})
        self.assertEqual(data['total'],0)
        self.assertTrue(all(d['accrued']==0 for d in data['distribution']))
    def test_effective_course_rule_overrides_default(self):
        rule=FinancialRule.objects.create(name='Само курс',course=self.course,valid_from=date(2027,1,1),valid_until=date(2027,1,31),
            university=30,faculty=15,center=0,confirmed=True,accounting_reference='Изключение')
        p=self.payment('100'); confirm_payment(p.pk,self.user); p.refresh_from_db()
        self.assertEqual(p.rule_id,rule.pk); self.assertEqual(p.allocation['remaining'],'55.00')
        duplicate=FinancialRule(name='Застъпване',valid_from=date(2027,1,1),university=30,faculty=10,center=10,confirmed=True,accounting_reference='тест')
        with self.assertRaises(ValidationError): duplicate.save()
    def test_unconfirmed_rules_do_not_apply(self):
        self.rule.confirmed=False; self.rule.save()
        p=self.payment()
        with self.assertRaises(ValidationError): confirm_payment(p.pk,self.user)
    def test_close_lock_carry_and_reopen_history(self):
        p=self.payment('100'); confirm_payment(p.pk,self.user)
        jan=Period.objects.create(title='Януари 2027',start=date(2027,1,1),end=date(2027,1,31))
        close_period(jan.pk,self.user,'approve','Утвърден доклад 1')
        close_period(jan.pk,self.user,'lock')
        feb=Period.objects.create(title='Февруари 2027',start=date(2027,2,1),end=date(2027,2,28))
        self.assertEqual(period_totals(feb)['opening'],'50.00')
        p2=self.payment('10')
        with self.assertRaises(ValidationError): confirm_payment(p2.pk,self.user)
        with self.assertRaises(ValidationError): close_period(jan.pk,self.user,'reopen')
        close_period(jan.pk,self.user,'reopen','Поправка на отчет 1')
        jan.refresh_from_db(); self.assertEqual(jan.revision,2)
        self.assertTrue(Audit.objects.filter(action='Отчет: reopen').exists())
    def test_expense_payment_is_separate_and_cannot_overpay(self):
        rate=CompensationRate.objects.create(name='Историческа ставка',category='external',percent=Decimal('11.86'),
            valid_from=date(2027,1,1),confirmed=True,accounting_reference='Счетоводство')
        exp=Expense.objects.create(date=date(2027,1,15),payee='Лектор',category='external',amount=Decimal('100'),rate=rate,state='approved',created_by=self.user)
        self.assertEqual(exp.total,Decimal('111.86')); self.assertEqual(exp.paid,0)
        payment=Disbursement(expense=exp,date=date(2027,1,16),amount=Decimal('50'),reference='expense-1')
        pay_expense(payment,self.user); self.assertEqual(exp.paid,Decimal('50'))
        with self.assertRaises(ValidationError):
            pay_expense(Disbursement(expense=exp,date=date(2027,1,17),amount=Decimal('70'),reference='expense-2'),self.user)
    def test_budget_break_even_uses_effective_rule(self):
        data=budget(self.course)
        self.assertEqual(data['revenue'],Decimal('900'))
        self.assertEqual(data['result'],Decimal('350'))
        self.assertEqual(data['break_even'],3)
    def test_budget_keeps_plan_after_expense_approval(self):
        exp=Expense.objects.create(date=date(2027,1,15),course=self.course,payee='Лектор',category='external',amount=Decimal('120'),state='planned',created_by=self.user)
        self.assertEqual(budget(self.course)['cost'],Decimal('120'))
        exp.amount=Decimal('140'); exp.state='approved'; exp.save()
        self.assertEqual(budget(self.course)['cost'],Decimal('120'))
        self.assertEqual(exp.total,Decimal('140'))
    def test_references_are_unique_across_financial_registers(self):
        from .models import Transfer
        self.payment(reference='shared-bank-reference')
        with self.assertRaises(ValidationError):
            Transfer.objects.create(date=date(2027,1,16),destination='university',amount=Decimal('10'),reference='shared-bank-reference',created_by=self.user)
    def test_viewer_cannot_write_and_csrf_required(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get('/new/payment/').status_code,403)
        p=self.payment()
        self.assertEqual(self.client.post(f'/payments/{p.pk}/confirm/').status_code,403)
        from django.test import Client
        c=Client(enforce_csrf_checks=True); c.force_login(self.user)
        self.assertEqual(c.post(f'/payments/{p.pk}/confirm/').status_code,403)
    def test_login_and_private_file_access(self):
        self.assertEqual(self.client.get('/').status_code,302)
        self.assertEqual(self.client.get('/private/payments/a.pdf').status_code,404)
    def test_course_date_payment_date_are_independent(self):
        self.course.start=timezone.make_aware(datetime(2025,1,1,10)); self.course.save()
        p=self.payment('90',kind='recording'); confirm_payment(p.pk,self.user)
        self.assertEqual(dashboard_data({'year':2027,'month':1})['total'],Decimal('90'))
        self.assertEqual(dashboard_data({'year':2025})['actual'],0)

class HistoricalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('import_history',verbosity=0)
        cls.user=get_user_model().objects.create_superuser('admin','',None)
    def test_all_fourteen_sources_and_no_fictional_transactions(self):
        self.assertEqual(SourceDocument.objects.count(),14)
        self.assertEqual(Period.objects.filter(historical=True).count(),14)
        self.assertEqual(Participant.objects.count(),0); self.assertEqual(Payment.objects.count(),0)
        self.assertEqual(HistoricalValue.objects.count(),225)
        self.assertTrue(Course.objects.get(title__startswith='Лидери за промяна').lecturers.filter(name__contains='Манол').exists())
    def test_idempotent_import(self):
        call_command('import_history',verbosity=0)
        self.assertEqual(SourceDocument.objects.count(),14); self.assertEqual(HistoricalValue.objects.count(),225)
    def test_real_source_acceptance_values(self):
        expected={
            'Юни 2026':['7146.00','2143.80','714.60','714.60','3573.00','505.43','4078.43','3534.78','543.65'],
            'Юли 2026':['2412.00','723.60','241.20','241.20','1206.00','543.65','1749.65','1715.07','34.58'],
            'Септември–октомври 2026':['2880.00','864.00','288.00','288.00','1440.00','34.58','1474.58','719.26','755.32']}
        keys=['total','university','faculty','center','remaining','opening','available','costs','closing']
        for title,values in expected.items():
            period=Period.objects.get(title=title); data=period_totals(period)
            self.assertEqual([Decimal(data[k]) for k in keys],list(map(Decimal,values)))
    def test_july_error_preserved_and_detected(self):
        p=Period.objects.get(title='Юли 2026'); row=p.values.get(row=0)
        self.assertEqual(row.original,Decimal('432')); self.assertEqual(row.recomputed_original,Decimal('336'))
        self.assertTrue(p.discrepancies.filter(code='fees_0',proposed='336.00').exists())
        self.assertEqual(p.state,'draft')
    def test_march_and_december_discrepancies(self):
        self.assertTrue(Discrepancy.objects.filter(period__title='Февруари–март 2026',code='arithmetic_costs',proposed='7724.34').exists())
        self.assertEqual(Discrepancy.objects.filter(period__title='Декември 2025',proposed='340.00').count(),2)
        self.assertTrue(Discrepancy.objects.filter(code='carry_forward').exists())
    def test_original_graphics_and_table_geometry_retained(self):
        from .documents import render_section,officials_for
        p=Period.objects.get(title='Юли 2026'); path=settings.BASE_DIR/SourceDocument.objects.get(path__contains='/01_').path
        d=render_section(path,'plan',p,period_totals(p),officials_for(p,date(2026,7,14)),date(2026,7,14))
        out=BytesIO(); d.save(out)
        with ZipFile(path) as original,ZipFile(out) as generated:
            for name in original.namelist():
                if name.startswith('word/media/'):
                    self.assertEqual(original.read(name),generated.read(name))
        base=Document(path)
        self.assertEqual(d.tables[0]._tbl.tblPr.xml,base.tables[0]._tbl.tblPr.xml)
        self.assertEqual(d.tables[0]._tbl.tblGrid.xml,base.tables[0]._tbl.tblGrid.xml)
        self.assertEqual(len(d.tables[0].rows),18)
        self.assertAlmostEqual(d.sections[0].page_width.mm,210,places=1)
        self.assertAlmostEqual(d.sections[0].page_height.mm,297,places=1)
        self.assertIn('432,00 евро',d.tables[0].cell(0,1).text)
    def test_historical_text_correction_requires_explicit_decision(self):
        from .documents import render_section,officials_for,txt
        p=Period.objects.get(title='Юли 2026'); day=date(2026,7,14)
        template=settings.BASE_DIR/p.source.path
        doc=render_section(template,'dean_report',p,period_totals(p),officials_for(p,day),day)
        original='РЪКОВОДИТЕЛ НА КУРСОВЕ КЪМ ПЕДАГОГИЧЕСКИЯ ФАКУЛТЕТ'
        self.assertIn(original,[x.text for x in doc.paragraphs])
        p.discrepancies.filter(code='dean_signature_role').update(state='accepted')
        revised=render_section(template,'dean_report',p,period_totals(p),officials_for(p,day),day)
        self.assertNotIn(original,[x.text for x in revised.paragraphs])
    def test_dynamic_orders_and_pdf_unavailable_are_honest(self):
        from .documents import generate
        p=Period.objects.get(title='Септември–октомври 2026')
        with tempfile.TemporaryDirectory() as folder,override_settings(MEDIA_ROOT=Path(folder),SOFFICE='missing-soffice'):
            d=generate(p,self.user,date(2026,10,14))
            self.assertIn('external_order.docx',d.files)
            self.assertNotIn('staff_order.docx',d.files)
            self.assertNotIn('invoice_order.docx',d.files)
            self.assertTrue(d.data['pdf_errors']); self.assertEqual(d.status,'draft')
            self.assertEqual(Payment.objects.count(),0)
    def test_adjustments_preserve_original(self):
        self.client.force_login(self.user)
        row=Period.objects.get(title='Юли 2026').values.get(row=0)
        response=self.client.post(f'/history/{row.pk}/adjust/',{'amount':'336.00','reason':'Потвърдена проверка на таксите'})
        self.assertEqual(response.status_code,302); row.refresh_from_db()
        self.assertEqual(row.original,Decimal('432')); self.assertEqual(row.effective_eur,Decimal('336'))
        self.assertTrue(HistoricalAdjustment.objects.filter(row=row).exists())
    def test_all_main_screens_render(self):
        self.client.force_login(self.user)
        paths=['/','/courses/',f'/courses/{Course.objects.first().pk}/','/register/payments/','/register/expenses/','/register/periods/',
               '/register/discrepancies/','/register/documents/','/register/participants/',f'/periods/{Period.objects.first().pk}/','/google/',
               '/new/course/','/new/payment/','/new/period/','/export/xlsx/','/export/csv/']
        for path in paths: self.assertEqual(self.client.get(path).status_code,200,path)

class GoogleTests(TestCase):
    def setUp(self):
        self.user=get_user_model().objects.create_superuser('admin','',None); self.client.force_login(self.user)
    @override_settings(GOOGLE_CLIENT_ID='',GOOGLE_CLIENT_SECRET='',TOKEN_ENCRYPTION_KEY='')
    def test_missing_configuration_never_simulates_connection(self):
        from .google import configuration
        with self.assertRaises(ValidationError): configuration()
        self.assertRedirects(self.client.post('/google/connect/'),'/google/')
    @patch('ledger.google.requests.post')
    def test_invalid_oauth_state_does_not_exchange_token(self,post):
        response=self.client.get('/google/callback/?state=forged&code=fake')
        self.assertEqual(response.status_code,302); post.assert_not_called()
    @patch('ledger.google.call')
    def test_response_sync_deduplicates_without_payments(self,call):
        from .google import sync_responses
        from .models import PriceCategory
        course=Course.objects.create(title='Google тест',form_id='real-form',google_question_map={
            'owner':self.user.pk,'name':'q1','email':'q2','phone':'q3','price':'q4','invoice':'q5','prices':{}})
        price=PriceCategory.objects.create(course=course,name='Стандартна',amount=Decimal('90'))
        course.google_question_map['prices']={'Стандартна':price.pk}; course.save()
        answers={key:{'textAnswers':{'answers':[{'value':value}]}} for key,value in [('q1','Име Фамилия'),('q2','test@example.com'),('q3','0880000000'),('q4','Стандартна')]}
        call.return_value={'responses':[{'responseId':'r1','answers':answers}]}
        self.assertEqual(sync_responses(course,self.user),(1,0))
        self.assertEqual(sync_responses(course,self.user),(0,0))
        self.assertEqual(Participant.objects.count(),1); self.assertEqual(Payment.objects.count(),0)
